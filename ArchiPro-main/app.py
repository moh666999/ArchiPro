from flask import (
    Flask,
    render_template,
    request,
    redirect,
    session,
    send_file,
    flash
)

import sqlite3
import os
import shutil
import zipfile

from datetime import datetime

from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename

import libsql_client
from dotenv import load_dotenv

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
env_path = os.path.join(BASE_DIR, "ArchiPro.env")
load_dotenv(dotenv_path=env_path)
load_dotenv(dotenv_path=env_path)
print("DEBUG URL:", os.getenv("TURSO_DATABASE_URL"))
print("DEBUG TOKEN:", os.getenv("TURSO_AUTH_TOKEN")[:10] if os.getenv("TURSO_AUTH_TOKEN") else "None")
# =====================================================
# APP
# =====================================================

app = Flask(__name__)

app.secret_key = "ArchiPro_secret_key"


# =====================================================
# PATHS
# =====================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DB_PATH = os.path.join(
    BASE_DIR,
    "database.db"
)

BACKUP_DIR = os.path.join(
    BASE_DIR,
    "backups"
)

ARCHIVE_FILES_DIR = os.path.join(
    BASE_DIR,
    "Archive_Files"
)

os.makedirs(
    ARCHIVE_FILES_DIR,
    exist_ok=True
)

os.makedirs(
    BACKUP_DIR,
    exist_ok=True
)


# =====================================================
# ARCHIPRO OPTIONS
# =====================================================

ARCHIVE_TYPES = [
    "ورقي",
    "الكتروني",
    "ورقي الكتروني"
]

DOCUMENT_TYPES = [
    "عادي",
    "سري",
    "سري للغاية"
]

def get_turso_db():
    url = os.getenv("TURSO_DATABASE_URL")
    auth_token = os.getenv("TURSO_AUTH_TOKEN")
    
    if not url:
        raise ValueError("Turso Database URL is missing inside get_turso_db!")
        
    return libsql_client.create_client_sync(url=url, auth_token=auth_token)

# =====================================================
# DATABASE
# =====================================================

# =====================================================
# DATABASE (TURSO WRAPPER FOR EXISTING ROUTES)
# =====================================================

class DictRow:
    """Mimics sqlite3.Row dict-like and tuple-like behavior."""
    def __init__(self, row_tuple, columns):
        self._tuple = row_tuple
        self._dict = {columns[i]: row_tuple[i] for i in range(len(columns))}

    def __getitem__(self, key):
        if isinstance(key, int):
            return self._tuple[key]
        return self._dict[key]

    def keys(self):
        return self._dict.keys()

class TursoCursor:
    """Mimics sqlite3 cursor/connection execution results."""
    def __init__(self, result):
        self.rows = []
        if result and hasattr(result, 'rows') and result.rows:
            for row in result.rows:
                self.rows.append(DictRow(tuple(row), result.columns))

    def fetchone(self):
        return self.rows[0] if self.rows else None

    def fetchall(self):
        return self.rows

class TursoDatabaseConnection:
    """Wraps libsql_client so it acts just like a sqlite3 connection object."""
    def __init__(self):
        self.client = get_turso_db()

    def execute(self, sql, params=()):
        try:
            # Clean up sqlite-specific PRAGMA commands if any slip through
            if "PRAGMA table_info" in sql:
                # Return an empty result structure for migrations so they don't crash
                class EmptyResult:
                    rows = []
                    columns = []
                return TursoCursor(EmptyResult())
            
            result = self.client.execute(sql, params)
            return TursoCursor(result)
        except Exception as e:
            error_str = str(e)
            if "UNIQUE constraint failed" in error_str or "already exists" in error_str:
                raise sqlite3.IntegrityError(error_str)
            raise e

    def commit(self):
        pass  # Turso auto-commits

    def close(self):
        self.client.close()

def get_db():
    return TursoDatabaseConnection()


# =====================================================
# LOGIN CHECK
# =====================================================

def check_login(username, password):
    conn = get_db()
    
    user = conn.execute(
        """
        SELECT *
        FROM users
        WHERE username = ? AND password = ?
        """,
        (username, password)
    ).fetchone()
    
    conn.close()
    return user


# =====================================================
# CURRENT USER
# =====================================================

def get_current_user():

    if "user_id" not in session:
        return None

    conn = get_db()

    user = conn.execute(
        """
        SELECT *
        FROM users
        WHERE id=?
        """,
        (
            session["user_id"],
        )
    ).fetchone()

    conn.close()

    return user


# =====================================================
# GENERAL PERMISSIONS
# =====================================================

def has_permission(permission):

    user = get_current_user()

    if not user:
        return False

    # Main User has everything
    if user["role"] == "main_user":
        return True

    try:

        return bool(
            user[permission]
        )

    except (KeyError, IndexError):

        return False


# =====================================================
# DOCUMENT TYPE PERMISSIONS
# =====================================================

def has_document_type_permission(document_type):

    user = get_current_user()

    if not user:
        return False

    # Main User can access all document types
    if user["role"] == "main_user":
        return True

    permission_map = {
        "عادي": "can_normal",
        "سري": "can_secret",
        "سري للغاية": "can_top_secret"
    }

    permission = permission_map.get(
        document_type
    )

    if not permission:
        return False

    try:

        return bool(
            user[permission]
        )

    except (KeyError, IndexError):

        return False


# =====================================================
# FILTER FILES BY DOCUMENT TYPE
# =====================================================

def filter_files_by_document_permission(files):

    user = get_current_user()

    if not user:
        return []

    # Main User sees everything
    if user["role"] == "main_user":
        return files

    allowed_files = []

    for file in files:

        document_type = (
            file["document_type"]
            if (
                "document_type" in file.keys()
                and file["document_type"]
            )
            else ""
        )

        if has_document_type_permission(
            document_type
        ):

            allowed_files.append(file)

    return allowed_files


# =====================================================
# HISTORY
# =====================================================

def add_history(file_number, action):

    user = get_current_user()

    username = (
        user["username"]
        if user
        else ""
    )

    action_date = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    device_name = request.headers.get(
        "User-Agent",
        "Unknown Device"
    )[:250]

    conn = get_db()

    conn.execute(
        """
        INSERT INTO history(
            file_number,
            action,
            action_date,
            username,
            device_name
        )
        VALUES(?,?,?,?,?)
        """,
        (
            file_number,
            action,
            action_date,
            username,
            device_name
        )
    )

    conn.commit()

    conn.close()


# =====================================================
# LOGIN
# =====================================================

@app.route(
    "/",
    methods=["GET", "POST"]
)
def login():

    message = ""

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        try:

            user = check_login(
                username,
                password
            )

        except Exception as e:

            return (
                f"Login error: {type(e).__name__}: {e}",
                500
            )

        if user:

            session["user_id"] = user["id"]

            session["username"] = user["username"]

            session["role"] = user["role"]

            return redirect("/dashboard")

        message = "Invalid username or password"

    return render_template(
        "login.html",
        message=message,
        error=message
    )


# =====================================================
# DASHBOARD
# =====================================================

@app.route("/dashboard")
def dashboard():

    user = get_current_user()

    if not user:
        return redirect("/")

    conn = get_db()

    files = conn.execute(
        """
        SELECT *
        FROM archive
        ORDER BY id DESC
        """
    ).fetchall()

    conn.close()

    # ---------------------------------------------
    # VIEW PERMISSION
    # ---------------------------------------------

    if not has_permission("can_view"):

        files = []

    else:

        files = filter_files_by_document_permission(
            files
        )

    return render_template(
        "dashboard.html",
        username=user["username"],
        user=user,
        files=files,
        search_text=""
    )


# =====================================================
# SEARCH
# =====================================================

@app.route("/search")
def search():

    user = get_current_user()

    if not user:
        return redirect("/")

    if not has_permission("can_view"):

        flash(
            "You do not have permission to view files."
        )

        return redirect("/dashboard")

    text = request.args.get(
        "q",
        ""
    ).strip()

    conn = get_db()

    if text:

        search_text = "%" + text + "%"

        files = conn.execute(
            """
            SELECT *
            FROM archive
            WHERE
                CAST(file_number AS TEXT) LIKE ?
                OR COALESCE(file_name, '') LIKE ?
                OR COALESCE(subject, '') LIKE ?
                OR COALESCE(department, '') LIKE ?
                OR COALESCE(document_type, '') LIKE ?
                OR COALESCE(archive_type, '') LIKE ?
                OR COALESCE(document_date, '') LIKE ?
                OR COALESCE(file_path, '') LIKE ?
                OR COALESCE(document_file_path, '') LIKE ?
                OR COALESCE(status, '') LIKE ?
                OR COALESCE(created_at, '') LIKE ?
                OR COALESCE(updated_at, '') LIKE ?
            ORDER BY id DESC
            """,
            (
                search_text,
                search_text,
                search_text,
                search_text,
                search_text,
                search_text,
                search_text,
                search_text,
                search_text,
                search_text,
                search_text,
                search_text
            )
        ).fetchall()

    else:

        files = conn.execute(
            """
            SELECT *
            FROM archive
            ORDER BY id DESC
            """
        ).fetchall()

    conn.close()

    # ---------------------------------------------
    # DOCUMENT TYPE PERMISSION FILTER
    # ---------------------------------------------

    files = filter_files_by_document_permission(
        files
    )

    return render_template(
        "dashboard.html",
        username=user["username"],
        user=user,
        files=files,
        search_text=text
    )


# =====================================================
# NEXT FILE NUMBER
# =====================================================

def get_next_file_number(conn):

    rows = conn.execute(
        """
        SELECT file_number
        FROM archive
        """
    ).fetchall()

    numbers = []

    for row in rows:

        value = str(
            row["file_number"] or ""
        ).strip()

        if value.isdigit():

            numbers.append(
                int(value)
            )

    next_number = max(
        numbers,
        default=0
    ) + 1

    return f"{next_number:03d}"


# =====================================================
# ADD FILE
# =====================================================

@app.route(
    "/add-file",
    methods=["GET", "POST"]
)
def add_file():

    if "user_id" not in session:
        return redirect("/")

    if not has_permission("can_add"):

        flash(
            "You do not have permission to add files."
        )

        return redirect("/dashboard")

    if request.method == "POST":

        file_name = request.form.get(
            "file_name",
            ""
        ).strip()

        subject = request.form.get(
            "subject",
            ""
        ).strip()

        department = request.form.get(
            "department",
            ""
        ).strip()

        document_type = request.form.get(
            "document_type",
            ""
        ).strip()

        archive_type = request.form.get(
            "archive_type",
            ""
        ).strip()

        document_date = request.form.get(
            "document_date",
            ""
        ).strip()

        file_path = request.form.get(
            "file_path",
            ""
        ).strip()

        status = request.form.get(
            "status",
            "Active"
        ).strip()

        document_file = request.files.get(
            "document_file"
        )

        # ---------------------------------------------
        # REQUIRED FIELDS
        # ---------------------------------------------

        if not file_name or not subject:

            flash(
                "Please enter the file name and subject."
            )

            return render_template(
                "add_file.html",
                archive_types=ARCHIVE_TYPES,
                document_types=DOCUMENT_TYPES
            )

        # ---------------------------------------------
        # VALIDATE ARCHIVE TYPE
        # ---------------------------------------------

        if archive_type not in ARCHIVE_TYPES:

            flash(
                "Please select a valid archive type."
            )

            return render_template(
                "add_file.html",
                archive_types=ARCHIVE_TYPES,
                document_types=DOCUMENT_TYPES
            )

        # ---------------------------------------------
        # VALIDATE DOCUMENT TYPE
        # ---------------------------------------------

        if document_type not in DOCUMENT_TYPES:

            flash(
                "Please select a valid document type."
            )

            return render_template(
                "add_file.html",
                archive_types=ARCHIVE_TYPES,
                document_types=DOCUMENT_TYPES
            )

        # ---------------------------------------------
        # DOCUMENT TYPE PERMISSION
        # ---------------------------------------------

        if not has_document_type_permission(
            document_type
        ):

            flash(
                "You do not have permission for this document type."
            )

            return render_template(
                "add_file.html",
                archive_types=ARCHIVE_TYPES,
                document_types=DOCUMENT_TYPES
            )

        now = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        conn = get_db()

        saved_document_path = ""

        try:

            # -----------------------------------------
            # AUTOMATIC FILE NUMBER
            # -----------------------------------------

            file_number = get_next_file_number(
                conn
            )

            # -----------------------------------------
            # SAVE ELECTRONIC DOCUMENT
            # -----------------------------------------

            if (
                document_file
                and document_file.filename
            ):

                safe_name = secure_filename(
                    document_file.filename
                )

                if not safe_name:

                    flash(
                        "The selected document has an invalid file name."
                    )

                    conn.close()

                    return render_template(
                        "add_file.html",
                        archive_types=ARCHIVE_TYPES,
                        document_types=DOCUMENT_TYPES
                    )

                target_name = (
                    f"{file_number}_{safe_name}"
                )

                target_path = os.path.join(
                    ARCHIVE_FILES_DIR,
                    target_name
                )

                document_file.save(
                    target_path
                )

                saved_document_path = os.path.join(
                    "Archive_Files",
                    target_name
                )

            # -----------------------------------------
            # INSERT FILE
            # -----------------------------------------

            conn.execute(
                """
                INSERT INTO archive(
                    file_number,
                    file_name,
                    subject,
                    department,
                    document_type,
                    archive_type,
                    document_date,
                    file_path,
                    document_file_path,
                    created_at,
                    updated_at,
                    status
                )
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    file_number,
                    file_name,
                    subject,
                    department,
                    document_type,
                    archive_type,
                    document_date,
                    file_path,
                    saved_document_path,
                    now,
                    now,
                    status
                )
            )

            conn.commit()

            add_history(
                file_number,
                "Added File"
            )

            flash(
                f"File {file_number} added successfully."
            )

            return redirect(
                "/dashboard"
            )

        except sqlite3.IntegrityError:

            flash(
                "Could not create the file number. Please try again."
            )

        finally:

            conn.close()

    return render_template(
        "add_file.html",
        archive_types=ARCHIVE_TYPES,
        document_types=DOCUMENT_TYPES
    )


# =====================================================
# EDIT FILE
# =====================================================

@app.route(
    "/edit-file/<int:file_id>",
    methods=["GET", "POST"]
)
def edit_file(file_id):

    if "user_id" not in session:
        return redirect("/")

    if not has_permission("can_edit"):

        flash(
            "You do not have permission to edit files."
        )

        return redirect("/dashboard")

    conn = get_db()

    file = conn.execute(
        """
        SELECT *
        FROM archive
        WHERE id=?
        """,
        (file_id,)
    ).fetchone()

    if not file:

        conn.close()

        flash(
            "File not found."
        )

        return redirect("/dashboard")

    # ---------------------------------------------
    # CURRENT DOCUMENT TYPE PERMISSION
    # ---------------------------------------------

    current_document_type = (
        file["document_type"]
        if (
            "document_type" in file.keys()
            and file["document_type"]
        )
        else ""
    )

    if not has_document_type_permission(
        current_document_type
    ):

        conn.close()

        flash(
            "You do not have permission to edit this document type."
        )

        return redirect("/dashboard")

    if request.method == "POST":

        file_name = request.form.get(
            "file_name",
            ""
        ).strip()

        subject = request.form.get(
            "subject",
            ""
        ).strip()

        department = request.form.get(
            "department",
            ""
        ).strip()

        document_type = request.form.get(
            "document_type",
            ""
        ).strip()

        archive_type = request.form.get(
            "archive_type",
            ""
        ).strip()

        document_date = request.form.get(
            "document_date",
            ""
        ).strip()

        file_path = request.form.get(
            "file_path",
            ""
        ).strip()

        status = request.form.get(
            "status",
            file["status"] or "Active"
        ).strip()

        document_file = request.files.get(
            "document_file"
        )

        # ---------------------------------------------
        # REQUIRED
        # ---------------------------------------------

        if not file_name or not subject:

            flash(
                "Please enter the file name and subject."
            )

            conn.close()

            return render_template(
                "edit_file.html",
                file=file,
                archive_types=ARCHIVE_TYPES,
                document_types=DOCUMENT_TYPES
            )

        # ---------------------------------------------
        # VALIDATE ARCHIVE TYPE
        # ---------------------------------------------

        if archive_type not in ARCHIVE_TYPES:

            flash(
                "Please select a valid archive type."
            )

            conn.close()

            return render_template(
                "edit_file.html",
                file=file,
                archive_types=ARCHIVE_TYPES,
                document_types=DOCUMENT_TYPES
            )

        # ---------------------------------------------
        # VALIDATE DOCUMENT TYPE
        # ---------------------------------------------

        if document_type not in DOCUMENT_TYPES:

            flash(
                "Please select a valid document type."
            )

            conn.close()

            return render_template(
                "edit_file.html",
                file=file,
                archive_types=ARCHIVE_TYPES,
                document_types=DOCUMENT_TYPES
            )

        # ---------------------------------------------
        # NEW DOCUMENT TYPE PERMISSION
        # ---------------------------------------------

        if not has_document_type_permission(
            document_type
        ):

            flash(
                "You do not have permission for this document type."
            )

            conn.close()

            return render_template(
                "edit_file.html",
                file=file,
                archive_types=ARCHIVE_TYPES,
                document_types=DOCUMENT_TYPES
            )

        updated_at = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        saved_document_path = (
            file["document_file_path"]
            if (
                "document_file_path" in file.keys()
                and file["document_file_path"]
            )
            else ""
        )

        try:

            # -----------------------------------------
            # NEW ELECTRONIC DOCUMENT
            # -----------------------------------------

            if (
                document_file
                and document_file.filename
            ):

                safe_name = secure_filename(
                    document_file.filename
                )

                if not safe_name:

                    flash(
                        "The selected document has an invalid file name."
                    )

                    conn.close()

                    return render_template(
                        "edit_file.html",
                        file=file,
                        archive_types=ARCHIVE_TYPES,
                        document_types=DOCUMENT_TYPES
                    )

                target_name = (
                    f"{file['file_number']}_{safe_name}"
                )

                target_path = os.path.join(
                    ARCHIVE_FILES_DIR,
                    target_name
                )

                document_file.save(
                    target_path
                )

                saved_document_path = os.path.join(
                    "Archive_Files",
                    target_name
                )

            # -----------------------------------------
            # UPDATE
            # -----------------------------------------

            conn.execute(
                """
                UPDATE archive
                SET
                    file_name=?,
                    subject=?,
                    department=?,
                    document_type=?,
                    archive_type=?,
                    document_date=?,
                    file_path=?,
                    document_file_path=?,
                    updated_at=?,
                    status=?
                WHERE id=?
                """,
                (
                    file_name,
                    subject,
                    department,
                    document_type,
                    archive_type,
                    document_date,
                    file_path,
                    saved_document_path,
                    updated_at,
                    status,
                    file_id
                )
            )

            conn.commit()

            add_history(
                file["file_number"],
                "Edited File"
            )

            flash(
                "File updated successfully."
            )

            conn.close()

            return redirect(
                "/dashboard"
            )

        except sqlite3.IntegrityError:

            flash(
                "Could not update the file."
            )

    conn.close()

    return render_template(
        "edit_file.html",
        file=file,
        archive_types=ARCHIVE_TYPES,
        document_types=DOCUMENT_TYPES
    )


# =====================================================
# DELETE FILE
# =====================================================

@app.route(
    "/delete-file/<int:file_id>",
    methods=["POST"]
)
def delete_file(file_id):

    if "user_id" not in session:
        return redirect("/")

    if not has_permission("can_delete"):

        flash(
            "You do not have permission to delete files."
        )

        return redirect("/dashboard")

    conn = get_db()

    file = conn.execute(
        """
        SELECT *
        FROM archive
        WHERE id=?
        """,
        (file_id,)
    ).fetchone()

    if file:

        # -----------------------------------------
        # DOCUMENT TYPE PERMISSION
        # -----------------------------------------

        if not has_document_type_permission(
            file["document_type"]
            if (
                "document_type" in file.keys()
                and file["document_type"]
            )
            else ""
        ):

            conn.close()

            flash(
                "You do not have permission to delete this document type."
            )

            return redirect("/dashboard")

        file_number = file["file_number"]

        conn.execute(
            """
            DELETE FROM archive
            WHERE id=?
            """,
            (file_id,)
        )

        conn.commit()

        conn.close()

        add_history(
            file_number,
            "Deleted File"
        )

        flash(
            "File deleted successfully."
        )

    else:

        conn.close()

        flash(
            "File not found."
        )

    return redirect(
        "/dashboard"
    )


# =====================================================
# DOCUMENT
# =====================================================

@app.route(
    "/document/<int:file_id>"
)
def document(file_id):

    if "user_id" not in session:
        return redirect("/")

    if not has_permission("can_view"):

        flash(
            "You do not have permission to view files."
        )

        return redirect("/dashboard")

    conn = get_db()

    file = conn.execute(
        """
        SELECT *
        FROM archive
        WHERE id=?
        """,
        (file_id,)
    ).fetchone()

    conn.close()

    if not file:

        flash(
            "File not found."
        )

        return redirect("/dashboard")

    # ---------------------------------------------
    # DOCUMENT TYPE PERMISSION
    # ---------------------------------------------

    if not has_document_type_permission(
        file["document_type"]
        if (
            "document_type" in file.keys()
            and file["document_type"]
        )
        else ""
    ):

        flash(
            "You do not have permission to view this document type."
        )

        return redirect("/dashboard")

    relative_path = (
        file["document_file_path"]
        if (
            "document_file_path" in file.keys()
            and file["document_file_path"]
        )
        else ""
    )

    if not relative_path:

        flash(
            "No electronic document is attached to this archive record."
        )

        return redirect(
            "/dashboard"
        )

    full_path = os.path.abspath(
        os.path.join(
            BASE_DIR,
            relative_path
        )
    )

    archive_root = os.path.abspath(
        ARCHIVE_FILES_DIR
    )

    if (
        not full_path.startswith(
            archive_root + os.sep
        )
        or not os.path.isfile(full_path)
    ):

        flash(
            "The attached document could not be found."
        )

        return redirect(
            "/dashboard"
        )

    add_history(
        file["file_number"],
        "Viewed File"
    )

    return send_file(
        full_path,
        as_attachment=False
    )


# =====================================================
# VIEW FILE
# =====================================================

@app.route(
    "/view-file/<int:file_id>"
)
def view_file(file_id):

    if "user_id" not in session:
        return redirect("/")

    if not has_permission("can_view"):

        flash(
            "You do not have permission to view files."
        )

        return redirect("/dashboard")

    conn = get_db()

    file = conn.execute(
        """
        SELECT *
        FROM archive
        WHERE id=?
        """,
        (file_id,)
    ).fetchone()

    conn.close()

    if not file:

        flash(
            "File not found."
        )

        return redirect(
            "/dashboard"
        )

    # ---------------------------------------------
    # DOCUMENT TYPE PERMISSION
    # ---------------------------------------------

    if not has_document_type_permission(
        file["document_type"]
        if (
            "document_type" in file.keys()
            and file["document_type"]
        )
        else ""
    ):

        flash(
            "You do not have permission to view this document type."
        )

        return redirect("/dashboard")

    add_history(
        file["file_number"],
        "Viewed File"
    )

    return render_template(
        "view_file.html",
        file=file
    )


# =====================================================
# DOWNLOAD FILE
# =====================================================

@app.route(
    "/download-file/<int:file_id>"
)
def download_file(file_id):

    if "user_id" not in session:
        return redirect("/")

    if not has_permission("can_download"):

        flash(
            "You do not have permission to download files."
        )

        return redirect("/dashboard")

    conn = get_db()

    file = conn.execute(
        """
        SELECT *
        FROM archive
        WHERE id=?
        """,
        (file_id,)
    ).fetchone()

    conn.close()

    if not file:

        flash(
            "File not found."
        )

        return redirect(
            "/dashboard"
        )

    # ---------------------------------------------
    # DOCUMENT TYPE PERMISSION
    # ---------------------------------------------

    if not has_document_type_permission(
        file["document_type"]
        if (
            "document_type" in file.keys()
            and file["document_type"]
        )
        else ""
    ):

        flash(
            "You do not have permission to download this document type."
        )

        return redirect("/dashboard")

    file_path = (
        file["document_file_path"]
        if (
            "document_file_path" in file.keys()
            and file["document_file_path"]
        )
        else file["file_path"]
    )

    if not file_path:

        flash(
            "The stored file could not be found on the server."
        )

        return redirect(
            "/dashboard"
        )

    if os.path.isabs(file_path):

        full_path = os.path.abspath(
            file_path
        )

    else:

        full_path = os.path.abspath(
            os.path.join(
                BASE_DIR,
                file_path
            )
        )

    if not os.path.isfile(full_path):

        flash(
            "The stored file could not be found on the server."
        )

        return redirect(
            "/dashboard"
        )

    add_history(
        file["file_number"],
        "Downloaded File"
    )

    return send_file(
        full_path,
        as_attachment=True
    )


# =====================================================
# USERS MANAGEMENT
# =====================================================

@app.route("/users")
def users():

    if "user_id" not in session:
        return redirect("/")

    user = get_current_user()

    if user["role"] != "main_user":

        flash(
            "Only the Main User can manage users."
        )

        return redirect(
            "/dashboard"
        )

    conn = get_db()

    users_list = conn.execute(
        """
        SELECT
            id,
            username,
            role,
            can_view,
            can_download,
            can_add,
            can_edit,
            can_delete,
            can_normal,
            can_secret,
            can_top_secret
        FROM users
        ORDER BY username
        """
    ).fetchall()

    conn.close()

    return render_template(
        "users.html",
        users=users_list,
        username=user["username"]
    )


# =====================================================
# ADD USER
# =====================================================

@app.route(
    "/add-user",
    methods=["POST"]
)
def add_user():

    if "user_id" not in session:
        return redirect("/")

    user = get_current_user()

    if user["role"] != "main_user":

        flash(
            "Only the Main User can create users."
        )

        return redirect(
            "/dashboard"
        )

    username = request.form.get(
        "username",
        ""
    ).strip()

    password = request.form.get(
        "password",
        ""
    )

    # ---------------------------------------------
    # OLD PERMISSIONS
    # ---------------------------------------------

    can_view = (
        1 if request.form.get("can_view")
        else 0
    )

    can_download = (
        1 if request.form.get("can_download")
        else 0
    )

    can_add = (
        1 if request.form.get("can_add")
        else 0
    )

    can_edit = (
        1 if request.form.get("can_edit")
        else 0
    )

    can_delete = (
        1 if request.form.get("can_delete")
        else 0
    )

    # ---------------------------------------------
    # DOCUMENT TYPE PERMISSIONS
    # ---------------------------------------------

    can_normal = (
        1 if request.form.get("can_normal")
        else 0
    )

    can_secret = (
        1 if request.form.get("can_secret")
        else 0
    )

    can_top_secret = (
        1 if request.form.get("can_top_secret")
        else 0
    )

    if not username or not password:

        flash(
            "Username and password are required."
        )

        return redirect(
            "/users"
        )

    conn = get_db()

    try:

        conn.execute(
            """
            INSERT INTO users(
                username,
                password,
                role,
                can_view,
                can_download,
                can_add,
                can_edit,
                can_delete,
                can_normal,
                can_secret,
                can_top_secret
            )
            VALUES(?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                username,
                password,
                "normal_user",
                can_view,
                can_download,
                can_add,
                can_edit,
                can_delete,
                can_normal,
                can_secret,
                can_top_secret
            )
        )

        conn.commit()

        flash(
            "User created successfully."
        )

    except sqlite3.IntegrityError:

        flash(
            "Username already exists."
        )

    finally:

        conn.close()

    return redirect(
        "/users"
    )


# =====================================================
# UPDATE PERMISSIONS
# =====================================================

@app.route(
    "/update-permissions/<int:user_id>",
    methods=["POST"]
)
def update_permissions(user_id):

    if "user_id" not in session:
        return redirect("/")

    current_user = get_current_user()

    if current_user["role"] != "main_user":

        flash(
            "Only the Main User can change permissions."
        )

        return redirect(
            "/dashboard"
        )

    # ---------------------------------------------
    # OLD PERMISSIONS
    # ---------------------------------------------

    can_view = (
        1 if request.form.get("can_view")
        else 0
    )

    can_download = (
        1 if request.form.get("can_download")
        else 0
    )

    can_add = (
        1 if request.form.get("can_add")
        else 0
    )

    can_edit = (
        1 if request.form.get("can_edit")
        else 0
    )

    can_delete = (
        1 if request.form.get("can_delete")
        else 0
    )

    # ---------------------------------------------
    # DOCUMENT TYPE PERMISSIONS
    # ---------------------------------------------

    can_normal = (
        1 if request.form.get("can_normal")
        else 0
    )

    can_secret = (
        1 if request.form.get("can_secret")
        else 0
    )

    can_top_secret = (
        1 if request.form.get("can_top_secret")
        else 0
    )

    conn = get_db()

    conn.execute(
        """
        UPDATE users
        SET
            can_view=?,
            can_download=?,
            can_add=?,
            can_edit=?,
            can_delete=?,
            can_normal=?,
            can_secret=?,
            can_top_secret=?
        WHERE id=?
        """,
        (
            can_view,
            can_download,
            can_add,
            can_edit,
            can_delete,
            can_normal,
            can_secret,
            can_top_secret,
            user_id
        )
    )

    conn.commit()

    conn.close()

    flash(
        "Permissions updated successfully."
    )

    return redirect(
        "/users"
    )


# =====================================================
# DELETE USER
# =====================================================

@app.route(
    "/delete-user/<int:user_id>",
    methods=["POST"]
)
def delete_user(user_id):

    if "user_id" not in session:
        return redirect("/")

    current_user = get_current_user()

    if current_user["role"] != "main_user":

        flash(
            "Only the Main User can delete users."
        )

        return redirect(
            "/dashboard"
        )

    if user_id == current_user["id"]:

        flash(
            "The Main User cannot delete their own account."
        )

        return redirect(
            "/users"
        )

    conn = get_db()

    conn.execute(
        """
        DELETE FROM users
        WHERE id=?
        """,
        (user_id,)
    )

    conn.commit()

    conn.close()

    flash(
        "User deleted successfully."
    )

    return redirect(
        "/users"
    )


# =====================================================
# CHANGE PASSWORD
# =====================================================

@app.route(
    "/change-password",
    methods=["GET", "POST"]
)
def change_password():

    if "user_id" not in session:
        return redirect("/")

    user = get_current_user()

    if request.method == "POST":

        new_password = request.form.get(
            "new_password",
            ""
        )

        confirm_password = request.form.get(
            "confirm_password",
            ""
        )

        if not new_password:

            flash(
                "Password cannot be empty."
            )

        elif new_password != confirm_password:

            flash(
                "Passwords do not match."
            )

        else:

            conn = get_db()

            conn.execute(
                """
                UPDATE users
                SET password=?
                WHERE id=?
                """,
                (
                    new_password,
                    user["id"]
                )
            )

            conn.commit()

            conn.close()

            flash(
                "Password changed successfully."
            )

            return redirect(
                "/dashboard"
            )

    return render_template(
        "change_password.html",
        username=user["username"]
    )


# =====================================================
# HISTORY
# =====================================================

@app.route("/history")
def history():

    if "user_id" not in session:
        return redirect("/")

    user = get_current_user()

    if user["role"] != "main_user":

        flash(
            "Only the Main User can view history."
        )

        return redirect(
            "/dashboard"
        )

    conn = get_db()

    history_list = conn.execute(
        """
        SELECT
            id,
            file_number,
            action,
            action_date,
            username,
            device_name
        FROM history
        ORDER BY id DESC
        """
    ).fetchall()

    conn.close()

    return render_template(
        "history.html",
        history=history_list,
        username=user["username"]
    )


# =====================================================
# CLEAR HISTORY
# =====================================================

@app.route(
    "/clear-history",
    methods=["POST"]
)
def clear_history():

    if "user_id" not in session:
        return redirect("/")

    user = get_current_user()

    if user["role"] != "main_user":

        flash(
            "Only the Main User can clear history."
        )

        return redirect(
            "/dashboard"
        )

    conn = get_db()

    conn.execute(
        "DELETE FROM history"
    )

    conn.commit()

    conn.close()

    flash(
        "History cleared successfully."
    )

    return redirect(
        "/history"
    )


# =====================================================
# EXPORT EXCEL
# =====================================================

@app.route("/export-excel")
def export_excel():

    if "user_id" not in session:
        return redirect("/")

    user = get_current_user()

    if user["role"] != "main_user":

        flash(
            "Only the Main User can export Excel files."
        )

        return redirect(
            "/dashboard"
        )

    conn = get_db()

    files = conn.execute(
        """
        SELECT
            file_number,
            file_name,
            subject,
            department,
            document_type,
            archive_type,
            document_date,
            file_path,
            status
        FROM archive
        ORDER BY id DESC
        """
    ).fetchall()

    conn.close()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Archive"

    headers = [
        "File Number",
        "File Name",
        "Subject",
        "Department",
        "Document Type",
        "Archive Type",
        "Date",
        "Storage Location",
        "Status"
    ]

    sheet.append(headers)

    for file in files:

        sheet.append(
            [
                file["file_number"],
                file["file_name"] or "",
                file["subject"] or "",
                file["department"] or "",
                file["document_type"] or "",
                file["archive_type"] or "",
                file["document_date"] or "",
                file["file_path"] or "",
                file["status"] or ""
            ]
        )

    for column in sheet.columns:

        max_length = 0

        column_letter = (
            column[0].column_letter
        )

        for cell in column:

            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            max_length = max(
                max_length,
                len(value)
            )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            45
        )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    filename = (
        f"ArchiPro_Export_{timestamp}.xlsx"
    )

    export_path = os.path.join(
        BACKUP_DIR,
        filename
    )

    workbook.save(
        export_path
    )

    add_history(
        "",
        "Export Excel"
    )

    return send_file(
        export_path,
        as_attachment=True,
        download_name=filename
    )


# =====================================================
# IMPORT EXCEL
# =====================================================

@app.route(
    "/import-excel",
    methods=["POST"]
)
def import_excel():

    if "user_id" not in session:
        return redirect("/")

    user = get_current_user()

    if user["role"] != "main_user":

        flash(
            "Only the Main User can import Excel files."
        )

        return redirect(
            "/dashboard"
        )

    excel_file = request.files.get(
        "excel_file"
    )

    if (
        not excel_file
        or excel_file.filename == ""
    ):

        flash(
            "Please select an Excel file."
        )

        return redirect(
            "/dashboard"
        )

    filename = secure_filename(
        excel_file.filename
    )

    if not filename.lower().endswith(
        ".xlsx"
    ):

        flash(
            "Only .xlsx Excel files are allowed."
        )

        return redirect(
            "/dashboard"
        )

    temp_excel = os.path.join(
        BACKUP_DIR,
        f"import_temp_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"
    )

    excel_file.save(
        temp_excel
    )

    imported_count = 0
    skipped_count = 0

    try:

        workbook = load_workbook(
            temp_excel,
            read_only=True,
            data_only=True
        )

        sheet = workbook.active

        rows = list(
            sheet.iter_rows(
                values_only=True
            )
        )

        if not rows:

            flash(
                "The Excel file is empty."
            )

            workbook.close()

            os.remove(
                temp_excel
            )

            return redirect(
                "/dashboard"
            )

        headers = [
            str(value).strip()
            if value is not None
            else ""
            for value in rows[0]
        ]

        required_headers = [
            "File Number",
            "File Name",
            "Subject",
            "Department",
            "Document Type",
            "Archive Type",
            "Date",
            "Storage Location",
            "Status"
        ]

        if headers[:9] != required_headers:

            flash(
                "Invalid Excel format. Please use an Excel file exported from ArchiPro."
            )

            workbook.close()

            os.remove(
                temp_excel
            )

            return redirect(
                "/dashboard"
            )

        conn = get_db()

        try:

            for row in rows[1:]:

                values = list(
                    row[:9]
                )

                while len(values) < 9:
                    values.append("")

                def clean(value):

                    return (
                        str(value).strip()
                        if value is not None
                        else ""
                    )

                file_number = clean(
                    values[0]
                )

                file_name = clean(
                    values[1]
                )

                subject = clean(
                    values[2]
                )

                department = clean(
                    values[3]
                )

                document_type = clean(
                    values[4]
                )

                archive_type = clean(
                    values[5]
                )

                document_date = clean(
                    values[6]
                )

                file_path = clean(
                    values[7]
                )

                status = (
                    clean(values[8])
                    or "Active"
                )

                if not file_name or not subject:

                    skipped_count += 1

                    continue

                # -------------------------------------
                # Validate imported choices
                # -------------------------------------

                if (
                    document_type
                    and document_type not in DOCUMENT_TYPES
                ):

                    skipped_count += 1

                    continue

                if (
                    archive_type
                    and archive_type not in ARCHIVE_TYPES
                ):

                    skipped_count += 1

                    continue

                if not file_number:

                    file_number = get_next_file_number(
                        conn
                    )

                now = datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                )

                try:

                    conn.execute(
                        """
                        INSERT INTO archive(
                            file_number,
                            file_name,
                            subject,
                            department,
                            document_type,
                            archive_type,
                            document_date,
                            file_path,
                            document_file_path,
                            created_at,
                            updated_at,
                            status
                        )
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                        """,
                        (
                            file_number,
                            file_name,
                            subject,
                            department,
                            document_type,
                            archive_type,
                            document_date,
                            file_path,
                            "",
                            now,
                            now,
                            status
                        )
                    )

                    imported_count += 1

                except sqlite3.IntegrityError:

                    skipped_count += 1

                    continue

            conn.commit()

        finally:

            conn.close()

        workbook.close()

        os.remove(
            temp_excel
        )

        add_history(
            "",
            "Import Excel"
        )

        message = (
            f"{imported_count} file(s) imported successfully."
        )

        if skipped_count:

            message += (
                f" {skipped_count} row(s) skipped."
            )

        flash(
            message
        )

    except Exception as e:

        if os.path.exists(
            temp_excel
        ):

            os.remove(
                temp_excel
            )

        flash(
            f"Import failed: {e}"
        )

    return redirect(
        "/dashboard"
    )


# =====================================================
# BACKUP
# =====================================================

@app.route("/backup")
def backup():

    if "user_id" not in session:
        return redirect("/")

    user = get_current_user()

    if user["role"] != "main_user":

        flash(
            "Only the Main User can create backups."
        )

        return redirect(
            "/dashboard"
        )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    filename = (
        f"ArchiPro_Full_Backup_{timestamp}.zip"
    )

    backup_path = os.path.join(
        BACKUP_DIR,
        filename
    )

    conn = get_db()

    archive_files = conn.execute(
        """
        SELECT
            document_file_path,
            file_path
        FROM archive
        """
    ).fetchall()

    conn.close()

    added_files = set()

    with zipfile.ZipFile(
        backup_path,
        "w",
        zipfile.ZIP_DEFLATED
    ) as backup_zip:

        # ---------------------------------------------
        # DATABASE
        # ---------------------------------------------

        if os.path.isfile(DB_PATH):

            backup_zip.write(
                DB_PATH,
                "database.db"
            )

        # ---------------------------------------------
        # ARCHIVE FILES
        # ---------------------------------------------

        for row in archive_files:

            possible_paths = [
                row["document_file_path"],
                row["file_path"]
            ]

            for relative_path in possible_paths:

                if not relative_path:
                    continue

                if os.path.isabs(
                    relative_path
                ):

                    file_path = os.path.abspath(
                        relative_path
                    )

                else:

                    file_path = os.path.abspath(
                        os.path.join(
                            BASE_DIR,
                            relative_path
                        )
                    )

                if not os.path.isfile(
                    file_path
                ):

                    continue

                base_name = os.path.basename(
                    file_path
                )

                if not base_name:
                    continue

                archive_name = (
                    f"Archive_Files/{base_name}"
                )

                if archive_name in added_files:
                    continue

                backup_zip.write(
                    file_path,
                    archive_name
                )

                added_files.add(
                    archive_name
                )

    add_history(
        "",
        "Full Backup"
    )

    return send_file(
        backup_path,
        as_attachment=True,
        download_name=filename
    )


# =====================================================
# RESTORE
# =====================================================

@app.route(
    "/restore",
    methods=["GET", "POST"]
)
def restore():

    if "user_id" not in session:
        return redirect("/")

    user = get_current_user()

    if user["role"] != "main_user":

        flash(
            "Only the Main User can restore backups."
        )

        return redirect(
            "/dashboard"
        )

    if request.method == "POST":

        backup_file = request.files.get(
            "backup_file"
        )

        if (
            not backup_file
            or backup_file.filename == ""
        ):

            flash(
                "Please select a backup file."
            )

            return redirect(
                "/restore"
            )

        filename = secure_filename(
            backup_file.filename
        )

        if not (
            filename.lower().endswith(".zip")
            or filename.lower().endswith(".db")
        ):

            flash(
                "Only ArchiPro .zip or .db backup files are allowed."
            )

            return redirect(
                "/restore"
            )

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        temp_backup = os.path.join(
            BACKUP_DIR,
            f"restore_temp_{timestamp}"
        )

        os.makedirs(
            temp_backup,
            exist_ok=True
        )

        uploaded_backup = os.path.join(
            temp_backup,
            filename
        )

        backup_file.save(
            uploaded_backup
        )

        try:

            database_source = None
            archive_source = None

            # -----------------------------------------
            # ZIP
            # -----------------------------------------

            if filename.lower().endswith(".zip"):

                with zipfile.ZipFile(
                    uploaded_backup,
                    "r"
                ) as backup_zip:

                    names = backup_zip.namelist()

                    database_name = next(
                        (
                            name
                            for name in names
                            if name.replace(
                                "\\",
                                "/"
                            ).lower()
                            == "database.db"
                        ),
                        None
                    )

                    if not database_name:

                        raise ValueError(
                            "The backup ZIP does not contain database.db."
                        )

                    database_source = os.path.join(
                        temp_backup,
                        "database.db"
                    )

                    with backup_zip.open(
                        database_name
                    ) as source, open(
                        database_source,
                        "wb"
                    ) as target:

                        shutil.copyfileobj(
                            source,
                            target
                        )

                    archive_members = [
                        name
                        for name in names
                        if name.replace(
                            "\\",
                            "/"
                        ).startswith(
                            "Archive_Files/"
                        )
                    ]

                    if archive_members:

                        archive_source = os.path.join(
                            temp_backup,
                            "Archive_Files"
                        )

                        os.makedirs(
                            archive_source,
                            exist_ok=True
                        )

                        for member in archive_members:

                            safe_name = os.path.basename(
                                member.replace(
                                    "\\",
                                    "/"
                                )
                            )

                            if not safe_name:
                                continue

                            destination = os.path.join(
                                archive_source,
                                safe_name
                            )

                            with backup_zip.open(
                                member
                            ) as source, open(
                                destination,
                                "wb"
                            ) as target:

                                shutil.copyfileobj(
                                    source,
                                    target
                                )

            # -----------------------------------------
            # DB ONLY
            # -----------------------------------------

            else:

                database_source = uploaded_backup

            # -----------------------------------------
            # CHECK DATABASE
            # -----------------------------------------

            test_conn = sqlite3.connect(
                database_source
            )

            required_tables = {
                "users",
                "archive",
                "history"
            }

            existing_tables = {
                row[0]
                for row in test_conn.execute(
                    """
                    SELECT name
                    FROM sqlite_master
                    WHERE type='table'
                    """
                ).fetchall()
            }

            test_conn.close()

            missing_tables = (
                required_tables
                - existing_tables
            )

            if missing_tables:

                raise ValueError(
                    "The backup database is missing required tables: "
                    + ", ".join(
                        sorted(
                            missing_tables
                        )
                    )
                )

            # -----------------------------------------
            # BACKUP CURRENT DATABASE
            # -----------------------------------------

            before_restore = os.path.join(
                BACKUP_DIR,
                f"database_before_restore_{timestamp}.db"
            )

            if os.path.isfile(
                DB_PATH
            ):

                shutil.copy2(
                    DB_PATH,
                    before_restore
                )

            # -----------------------------------------
            # RESTORE DATABASE
            # -----------------------------------------

            shutil.copy2(
                database_source,
                DB_PATH
            )

            # -----------------------------------------
            # RESTORE ARCHIVE FILES
            # -----------------------------------------

            if (
                archive_source
                and os.path.isdir(
                    archive_source
                )
            ):

                restore_archive_dir = os.path.join(
                    BASE_DIR,
                    "Archive_Files"
                )

                os.makedirs(
                    restore_archive_dir,
                    exist_ok=True
                )

                for archive_filename in os.listdir(
                    archive_source
                ):

                    source_path = os.path.join(
                        archive_source,
                        archive_filename
                    )

                    if os.path.isfile(
                        source_path
                    ):

                        destination_path = os.path.join(
                            restore_archive_dir,
                            archive_filename
                        )

                        shutil.copy2(
                            source_path,
                            destination_path
                        )

            # -----------------------------------------
            # HISTORY
            # -----------------------------------------

            try:

                add_history(
                    "",
                    "Restore Backup"
                )

            except Exception:

                pass

            shutil.rmtree(
                temp_backup,
                ignore_errors=True
            )

            flash(
                "Backup restored successfully. Please log in again."
            )

            session.clear()

            return redirect("/")

        except Exception as e:

            shutil.rmtree(
                temp_backup,
                ignore_errors=True
            )

            flash(
                f"Restore failed: {e}"
            )

            return redirect(
                "/restore"
            )

    # ---------------------------------------------
    # AVAILABLE BACKUPS
    # ---------------------------------------------

    backups = []

    for filename in os.listdir(
        BACKUP_DIR
    ):

        if filename.lower().endswith(
            (".zip", ".db")
        ):

            backups.append(
                filename
            )

    backups.sort(
        reverse=True
    )

    return render_template(
        "restore.html",
        backups=backups,
        username=user["username"]
    )


# =====================================================
# LOGOUT
# =====================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect("/")


# =====================================================
# RUN
# =====================================================

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
